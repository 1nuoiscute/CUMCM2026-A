"""Compare all submitted/result cells with regenerated workbooks, ignoring ZIP metadata."""
from pathlib import Path
from numbers import Real
import json
import math
from itertools import zip_longest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

def compare(q):
    reference = ROOT / "results/submitted" / f"result{q}.xlsx"
    generated = ROOT / "outputs" / f"q{q}" / f"result{q}.xlsx"
    a = load_workbook(reference, data_only=True, read_only=True)
    b = load_workbook(generated, data_only=True, read_only=True)
    report = {"question": q, "sheet_names_equal": a.sheetnames == b.sheetnames,
              "numeric_cells": 0, "different_cells": 0, "max_absolute_difference": 0.0,
              "mismatches_over_tolerance": 0, "tolerance": 0.00005}
    try:
        if a.sheetnames != b.sheetnames:
            report["mismatches_over_tolerance"] += 1
            return report
        for name in a.sheetnames:
            for row_a, row_b in zip_longest(a[name].iter_rows(values_only=True),
                                           b[name].iter_rows(values_only=True), fillvalue=()):
                for x, y in zip_longest(row_a, row_b, fillvalue=None):
                    if isinstance(x, Real) and isinstance(y, Real):
                        report["numeric_cells"] += 1
                        delta = abs(x-y)
                        report["max_absolute_difference"] = max(report["max_absolute_difference"], delta)
                        if not (math.isfinite(x) and math.isfinite(y)) or delta > report["tolerance"]:
                            report["mismatches_over_tolerance"] += 1
                    elif x != y:
                        report["mismatches_over_tolerance"] += 1
                    if x != y:
                        report["different_cells"] += 1
        return report
    finally:
        a.close()
        b.close()

def main():
    reports = [compare(q) for q in range(1,5)]
    target = ROOT / "outputs/comparison.json"
    target.write_text(json.dumps(reports, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(reports, ensure_ascii=False, indent=2))
    raise SystemExit(1 if any(x["mismatches_over_tolerance"] for x in reports) else 0)

if __name__ == "__main__":
    main()
