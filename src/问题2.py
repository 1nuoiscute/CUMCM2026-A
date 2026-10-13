"""问题二：变物性条件下的圆柱药材干燥计算。"""
from pathlib import Path
import sys
from project_paths import DATA_ROOT, OUTPUT_ROOT

import matplotlib as mpl
mpl.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
SUPPORT = HERE
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

from drying_core import CylinderConfig, MaterialProperties, sample_radial_field, solve_fixed_radius

INPUT = DATA_ROOT / "附件1.xlsx"
FIG = OUTPUT_ROOT / "q2"
DATA = OUTPUT_ROOT / "q2"
RESULT = FIG / "result2.xlsx"


class Ambient:
    def __init__(self, path):
        ws = load_workbook(path, data_only=True, read_only=True).active
        a = np.asarray([r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None], float)
        self.t, self.T, self.C = a[:, 0], a[:, 1], a[:, 2]

    def __call__(self, t):
        time = np.asarray(t, dtype=float)
        ambient_temperature = np.interp(time, self.t, self.T)
        ambient_moisture = np.interp(time, self.t, self.C)
        ambient_temperature = np.where(time > 4.0 * 3600.0, 50.0, ambient_temperature)
        ambient_moisture = np.where(time > 4.0 * 3600.0, 0.05, ambient_moisture)
        return ambient_temperature, ambient_moisture


def rho(C, T): return 650.0 + 128.0 * C
def cp(C, T): return 1450.0 + 2736.0 * C / (C + 1.0)
def k(C, T): return 0.21 + 0.38 * C / (C + 1.0)
def D_full(C, T): return 2.4e-3 * np.exp(-0.45 / np.maximum(C, 1e-8)) * np.exp(-3850.0 / T)


def config():
    return CylinderConfig(0.02, 25.0, 8e-7, 301.15, 2.55,
        MaterialProperties(rho, cp, k, D_full))


def setup_plot():
    available = {f.name for f in mpl.font_manager.fontManager.ttflist}
    font = "Microsoft YaHei" if "Microsoft YaHei" in available else "SimHei"
    mpl.rcParams.update({"font.sans-serif": [font], "axes.unicode_minus": False,
                         "figure.dpi": 130, "savefig.dpi": 260,
                         "axes.spines.top": False, "axes.spines.right": False})


def save_workbook(t, radii_cm, temperature, moisture):
    wb = Workbook(); wb.remove(wb.active)
    for name, field in (("温度", temperature), ("水分浓度", moisture)):
        ws = wb.create_sheet(name)
        ws.append(["时间\\到药材中心的距离", *map(float, radii_cm)])
        for i, ti in enumerate(t): ws.append([int(ti), *[round(float(x), 4) for x in field[i]]])
        ws.freeze_panes = "B2"; ws.column_dimensions["A"].width = 24
    wb.save(RESULT)


def run_validation(ambient, baseline, radii_cm, target_times):
    """用1600区间细网格作参考，并用Radau核对时间积分。"""
    reference = solve_fixed_radius(config(), ambient, 10800.0, target_times,
        n_intervals=1600, method="BDF", max_step_s=5.0, rtol=1e-9, atol=1e-11)
    radau = solve_fixed_radius(config(), ambient, 10800.0, target_times,
        n_intervals=800, method="Radau", max_step_s=5.0, rtol=1e-9, atol=1e-11)
    refT, refC = sample_radial_field(reference, radii_cm/100)
    rT, rC = sample_radial_field(radau, radii_cm/100)
    bT_all, bC_all = sample_radial_field(baseline, radii_cm/100)
    base_indices = target_times.astype(int)
    bT, bC = bT_all[base_indices], bC_all[base_indices]
    rows = [
        ("BDF，800区间 vs. 1600区间", np.max(abs(bT-refT)), np.max(abs(bC-refC))),
        ("Radau vs. BDF（800区间）", np.max(abs(rT-bT)), np.max(abs(rC-bC))),
    ]
    def sci(value):
        mantissa, exponent = f"{value:.3e}".split("e")
        return "$" + mantissa + r"\times10^{" + str(int(exponent)) + "}$"
    lines=[f"{name} & {sci(dt)} & {sci(dc)} \\\\" for name,dt,dc in rows]
    (DATA/"validation_rows.tex").write_text("\n".join(lines)+"\n\\bottomrule\n",encoding="utf-8")
    return rows


def main():
    if not INPUT.is_file():
        raise FileNotFoundError(f"未找到附件1：{INPUT}")
    FIG.mkdir(parents=True, exist_ok=True); DATA.mkdir(parents=True, exist_ok=True); setup_plot()
    ambient = Ambient(INPUT)
    times = np.arange(0.0, 10800.0 + 1.0, 1.0)
    sol = solve_fixed_radius(config(), ambient, 10800.0, times,
                             n_intervals=800, method="BDF", max_step_s=5.0,
                             rtol=1e-9, atol=1e-11)
    radii_cm = np.round(np.arange(0, 2.01, 0.1), 1)
    Tout, Cout = sample_radial_field(sol, radii_cm / 100.0)
    save_workbook(times[1:], radii_cm, Tout[1:], Cout[1:])

    sample_h = np.arange(0.5, 3.01, 0.5)
    idx = (sample_h * 3600).astype(int)
    table_idx = [0, 5, 10, 15, 20]
    for fn, arr in (("temperature_rows.tex", Tout), ("moisture_rows.tex", Cout)):
        lines=[]
        for h, ii in zip(sample_h, idx):
            lines.append(f"{h:.1f} & " + " & ".join(f"{v:.4f}" for v in arr[ii, table_idx]) + r" \\")
        (DATA/fn).write_text("\n".join(lines)+"\n\\bottomrule\n", encoding="utf-8")

    # 图1：附件1是不断变化的外部驱动，不是常值。
    tq=np.linspace(0,10800,1081); Ta,Ca=ambient(tq)
    fig,ax=plt.subplots(2,1,figsize=(7.3,5.2),sharex=True,constrained_layout=True)
    ax[0].plot(tq/3600,Ta,color="#c2410c",lw=2); ax[0].set_ylabel("热风温度 / °C")
    ax[1].plot(tq/3600,Ca,color="#0369a1",lw=2); ax[1].set_ylabel("热风水分浓度 / (kg/kg)"); ax[1].set_xlabel("时间 / h")
    for a in ax:a.grid(alpha=.2)
    fig.savefig(FIG/"q2_boundary.png",bbox_inches="tight"); plt.close(fig)

    # 图2：内部分层响应。
    colors=plt.cm.viridis(np.linspace(.12,.9,len(idx)))
    fig,ax=plt.subplots(1,2,figsize=(9.2,3.7),constrained_layout=True)
    for c,h,ii in zip(colors,sample_h,idx):
        ax[0].plot(radii_cm,Tout[ii],color=c,label=f"{h:g} h")
        ax[1].plot(radii_cm,Cout[ii],color=c,label=f"{h:g} h")
    ax[0].set_ylabel("温度 / °C"); ax[1].set_ylabel("水分浓度 / (kg/kg)")
    for a in ax:a.set_xlabel("到中心距离 / cm"); a.grid(alpha=.2)
    ax[1].legend(frameon=False,ncol=2,fontsize=8)
    fig.savefig(FIG/"q2_profiles.png",bbox_inches="tight"); plt.close(fig)

    # 图3：将D拆成升温促进与失水抑制两个因子。
    T0=301.15; C0=2.55
    Dfield=D_full(Cout.T,Tout.T+273.15)
    fT=np.exp(-3850*(1/(Tout.T+273.15)-1/T0))
    fC=np.exp(-.45*(1/np.maximum(Cout.T,1e-8)-1/C0))
    fig,ax=plt.subplots(1,2,figsize=(9.2,3.7),constrained_layout=True)
    for c,h,ii in zip(colors,sample_h,idx):
        ax[0].plot(radii_cm,Dfield[:,ii]*1e9,color=c,label=f"{h:g} h")
    ax[0].set_xlabel("到中心距离 / cm"); ax[0].set_ylabel(r"$D$ / ($10^{-9}$ m$^2$/s)"); ax[0].grid(alpha=.2)
    ax[1].plot(radii_cm,fT[:,-1],color="#c2410c",lw=2,label="升温促进因子")
    ax[1].plot(radii_cm,fC[:,-1],color="#0369a1",lw=2,label="失水抑制因子")
    ax[1].axhline(1,color="#777",ls="--",lw=1); ax[1].set_xlabel("到中心距离 / cm"); ax[1].set_ylabel("相对于初始状态的倍率"); ax[1].legend(frameon=False); ax[1].grid(alpha=.2)
    fig.savefig(FIG/"q2_diffusivity.png",bbox_inches="tight"); plt.close(fig)

    summary = {
      "Tcenter": Tout[-1,0], "Tsurface": Tout[-1,-1],
      "Ccenter": Cout[-1,0], "Csurface": Cout[-1,-1],
      "Dcenter": Dfield[0,-1], "Dsurface": Dfield[-1,-1],
      "fTsurface": fT[-1,-1], "fCsurface": fC[-1,-1],
      "nfev": sol.nfev
    }
    D0=float(D_full(np.array([2.55]),np.array([301.15]))[0])
    summary.update({
      "Dzero": D0, "DnetSurface": Dfield[-1,-1]/D0,
      "DnetCenter": Dfield[0,-1]/D0,
      "fTcenter": fT[0,-1], "fCcenter": fC[0,-1],
      "deltaT": Tout[-1,-1]-Tout[-1,0],
      "deltaC": Cout[-1,0]-Cout[-1,-1],
      "BiInitial": 25*0.01/float(k(np.array([2.55]),np.array([301.15]))[0]),
      "BiSurface": 25*0.01/float(k(np.array([Cout[-1,-1]]),np.array([Tout[-1,-1]+273.15]))[0]),
    })
    validation = run_validation(ambient, sol, np.array([0,.5,1,1.5,2.0]), idx.astype(float))
    summary.update({"gridT":validation[0][1], "gridC":validation[0][2],
                    "radauT":validation[1][1], "radauC":validation[1][2]})
    commands=[]
    for name,value in summary.items():
        if name in ("Dcenter", "Dsurface", "Dzero", "gridT", "gridC", "radauT", "radauC"):
            mantissa, exponent = f"{value:.5e}".split("e")
            shown = mantissa + r"\times10^{" + str(int(exponent)) + "}"
        else:
            shown = f"{value:.6g}"
        commands.append(f"\\newcommand{{\\{name}}}{{{shown}}}")
    (DATA/"summary.tex").write_text("\n".join(commands)+"\n",encoding="utf-8")
    print(summary)

if __name__ == "__main__": main()
