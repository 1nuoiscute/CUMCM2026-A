"""A题问题三：沿用问题二模型，计算全域含水率达标时间。"""
from pathlib import Path
import shutil
import sys
from project_paths import DATA_ROOT, OUTPUT_ROOT

import matplotlib as mpl
mpl.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SUPPORT_ROOT = ROOT
CONTEST_ROOT = ROOT.parent
sys.path.insert(0, str(ROOT))

from drying_core import (
    CylinderConfig,
    MaterialProperties,
    radial_grid,
    sample_radial_field,
    solve_fixed_radius,
)


INPUT = DATA_ROOT / "附件1.xlsx"
TEMPLATE = DATA_ROOT / "附件3" / "result3.xlsx"
OUTPUT = OUTPUT_ROOT / "q3" / "result3.xlsx"
RESULTS = OUTPUT_ROOT / "q3"
FIGURES = RESULTS

R = 0.02
H = 25.0
HM = 8.0e-7
T0_K = 301.15
C0 = 2.55
THRESHOLD = 0.15
PLATEAU_START_S = 4.0 * 3600.0
PLATEAU_T_C = 50.0
PLATEAU_C = 0.0500


class FullProcessAmbient:
    """附件1分段线性边界；4 h 后采用稳定阶段均值。"""

    def __init__(self, path: Path):
        ws = load_workbook(path, data_only=True, read_only=True).active
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        values = np.asarray(rows, dtype=float)
        self.time_s = values[:, 0]
        self.temperature_c = values[:, 1]
        self.moisture = values[:, 2]
        if self.time_s[0] != 0 or self.time_s[-1] != PLATEAU_START_S:
            raise ValueError("附件1时间范围应为0~4 h")

    def __call__(self, time_s: float):
        if time_s <= PLATEAU_START_S:
            return (
                float(np.interp(time_s, self.time_s, self.temperature_c)),
                float(np.interp(time_s, self.time_s, self.moisture)),
            )
        return PLATEAU_T_C, PLATEAU_C


def rho(C, T):
    return 650.0 + 128.0 * C


def cp(C, T):
    return 1450.0 + 2736.0 * C / (C + 1.0)


def conductivity(C, T):
    return 0.21 + 0.38 * C / (C + 1.0)


def diffusivity(C, T):
    return 2.4e-3 * np.exp(-0.45 / np.maximum(C, 1.0e-8)) * np.exp(-3850.0 / T)


def model_config():
    return CylinderConfig(
        radius_m=R,
        heat_transfer_coefficient=H,
        mass_transfer_coefficient=HM,
        initial_temperature_k=T0_K,
        initial_moisture=C0,
        properties=MaterialProperties(rho, cp, conductivity, diffusivity),
    )


def interpolate_crossing(time_s, values, threshold=THRESHOLD):
    hit = np.flatnonzero(values <= threshold)
    if hit.size == 0:
        raise RuntimeError("72 h内未达到含水率阈值")
    i1 = int(hit[0])
    if i1 == 0:
        return float(time_s[0]), 0
    i0 = i1 - 1
    fraction = (values[i0] - threshold) / (values[i0] - values[i1])
    return float(time_s[i0] + fraction * (time_s[i1] - time_s[i0])), i1


def field_at_time(time_s, field, target_s):
    i1 = int(np.searchsorted(time_s, target_s, side="left"))
    if i1 == 0:
        return field[:, 0].copy()
    if i1 >= time_s.size:
        return field[:, -1].copy()
    if time_s[i1] == target_s:
        return field[:, i1].copy()
    i0 = i1 - 1
    w = (target_s - time_s[i0]) / (time_s[i1] - time_s[i0])
    return (1.0 - w) * field[:, i0] + w * field[:, i1]


def qualification_boundary(radius_m, moisture):
    rq = np.empty(moisture.shape[1])
    fraction = np.empty_like(rq)
    for j, profile in enumerate(moisture.T):
        if profile[-1] > THRESHOLD:
            rq[j], fraction[j] = R, 0.0
        elif profile[0] <= THRESHOLD:
            rq[j], fraction[j] = 0.0, 1.0
        else:
            i1 = int(np.flatnonzero(profile <= THRESHOLD)[0])
            i0 = i1 - 1
            w = (profile[i0] - THRESHOLD) / (profile[i0] - profile[i1])
            rq[j] = radius_m[i0] + w * (radius_m[i1] - radius_m[i0])
            fraction[j] = 1.0 - (rq[j] / R) ** 2
    return rq, fraction


def setup_plot():
    available = {f.name for f in mpl.font_manager.fontManager.ttflist}
    font = "Microsoft YaHei" if "Microsoft YaHei" in available else "SimHei"
    mpl.rcParams.update({
        "font.sans-serif": [font],
        "axes.unicode_minus": False,
        "figure.dpi": 130,
        "savefig.dpi": 320,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "font.size": 9,
    })


def save_figure(fig, name):
    fig.savefig(FIGURES / f"{name}.png", dpi=320, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def save_workbook(times_s, radii_cm, moisture):
    shutil.copy2(TEMPLATE, OUTPUT)
    wb = load_workbook(OUTPUT)
    ws = wb.active
    ws.delete_rows(1, ws.max_row)
    ws.append(["时间\\到药材中心的距离", *[round(float(x), 1) for x in radii_cm]])
    for i, time_s in enumerate(times_s):
        ws.append([int(round(time_s)), *[round(float(x), 4) for x in moisture[i]]])
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 24
    for column in ws.iter_cols(min_col=2, max_col=ws.max_column):
        ws.column_dimensions[column[0].column_letter].width = 10
    wb.save(OUTPUT)


def make_table_tex(hours, moisture_rows, completion_h):
    lines = []
    for h, row in zip(hours[:-1], moisture_rows[:-1]):
        lines.append(f"{int(h)} & " + " & ".join(f"{x:.4f}" for x in row) + r" \\")
    lines.append(
        f"烘干结束（{completion_h:.2f} h） & "
        + " & ".join(f"{x:.4f}" for x in moisture_rows[-1])
        + r" \\"
    )
    (RESULTS / "q3_table_rows.tex").write_text("\n".join(lines) + "\n\\bottomrule\n", encoding="utf-8")


def make_figures(solution, completion_s, completion_index, requested_radii_m, sampled_c):
    setup_plot()
    time_h = solution.time_s / 3600.0
    _, _, volumes = radial_grid(R, solution.radius_m.size - 1)
    weights = volumes / volumes.sum()
    mean_c = weights @ solution.moisture

    colors = ["#264653", "#2A9D8F", "#6A994E", "#E9C46A", "#E76F51"]
    fig, ax = plt.subplots(figsize=(6.5, 3.8), constrained_layout=True)
    for j, (r_cm, color) in enumerate(zip(requested_radii_m * 100.0, colors)):
        ax.plot(time_h[: completion_index + 1], sampled_c[: completion_index + 1, j],
                lw=1.4, color=color, label=f"r={r_cm:g} cm")
    ax.plot(time_h[: completion_index + 1], mean_c[: completion_index + 1],
            color="#555555", lw=1.2, ls="--", label="体积平均")
    ax.axhline(THRESHOLD, color="#B22222", lw=1.0, ls=":", label="达标阈值0.15")
    ax.axvline(completion_s / 3600.0, color="#B22222", lw=1.0, ls="--")
    ax.set(xlabel="时间 / h", ylabel="干基含水率 / (kg/kg)", xlim=(0, completion_s / 3600.0))
    ax.legend(ncol=3, frameon=False, fontsize=8)
    save_figure(fig, "result_q3_moisture_history")

    profile_hours = [6, 12, 24, 36, 48, completion_s / 3600.0]
    fig, ax = plt.subplots(figsize=(6.3, 3.8), constrained_layout=True)
    for h, color in zip(profile_hours, ["#440154", "#3B528B", "#21918C", "#5EC962", "#FDE725", "#D1495B"]):
        profile = field_at_time(solution.time_s, solution.moisture, h * 3600.0)
        label = f"{h:.2f} h" if h == profile_hours[-1] else f"{h:g} h"
        ax.plot(solution.radius_m * 100.0, profile, lw=1.5, color=color, label=label)
    ax.axhline(THRESHOLD, color="#B22222", lw=1.0, ls=":")
    ax.set(xlabel="到药材中心的距离 / cm", ylabel="干基含水率 / (kg/kg)", xlim=(0, 2.0))
    ax.legend(frameon=False, ncol=3, fontsize=8)
    save_figure(fig, "result_q3_radial_profiles")

def main():
    RESULTS.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    ambient = FullProcessAmbient(INPUT)
    output_times = np.arange(0.0, 72.0 * 3600.0 + 60.0, 60.0)
    solution = solve_fixed_radius(
        model_config(), ambient, 72.0 * 3600.0, output_times,
        n_intervals=800, method="BDF", max_step_s=5.0, rtol=1.0e-9, atol=1.0e-11,
    )

    max_c = solution.moisture.max(axis=0)
    completion_s, completion_index = interpolate_crossing(solution.time_s, max_c)

    output_radii_cm = np.arange(0.0, 2.0 + 0.1, 0.1)
    _, output_c = sample_radial_field(solution, output_radii_cm / 100.0)
    save_workbook(
        solution.time_s[1:completion_index + 1],
        output_radii_cm,
        output_c[1:completion_index + 1],
    )

    paper_radii_cm = np.array([0.0, 0.5, 1.0, 1.5, 2.0])
    _, paper_c = sample_radial_field(solution, paper_radii_cm / 100.0)
    six_hour_values = np.arange(6.0, np.floor(completion_s / 21600.0) * 6.0 + 0.1, 6.0)
    table_hours = np.append(six_hour_values, completion_s / 3600.0)
    table_rows = []
    for h in table_hours[:-1]:
        table_rows.append(field_at_time(solution.time_s, paper_c.T, h * 3600.0))
    completion_profile = field_at_time(solution.time_s, solution.moisture, completion_s)
    table_rows.append(np.interp(paper_radii_cm / 100.0, solution.radius_m, completion_profile))
    table_rows = np.asarray(table_rows)
    make_table_tex(table_hours, table_rows, completion_s / 3600.0)

    crossing_rows = []
    for j, r_cm in enumerate(paper_radii_cm):
        crossing_s, _ = interpolate_crossing(solution.time_s, paper_c[:, j])
        crossing_rows.append((r_cm, crossing_s / 3600.0))
    np.savetxt(
        RESULTS / "q3_threshold_times.csv",
        np.asarray(crossing_rows), delimiter=",", header="radius_cm,threshold_time_h", comments="",
        fmt=["%.1f", "%.8f"],
    )
    np.savetxt(
        RESULTS / "q3_table.csv",
        np.column_stack((table_hours, table_rows)), delimiter=",",
        header="time_h,C_0cm,C_0p5cm,C_1cm,C_1p5cm,C_2cm", comments="", fmt="%.8f",
    )

    make_figures(solution, completion_s, completion_index, paper_radii_cm / 100.0, paper_c)

    surface_hit_h = crossing_rows[-1][1]
    late_fraction = (completion_s / 3600.0 - surface_hit_h) / (completion_s / 3600.0)
    final_surface = table_rows[-1, -1]
    summary = (
        f"completion_time_h={completion_s/3600.0:.8f}\n"
        f"first_qualifying_output_minute_h={solution.time_s[completion_index]/3600.0:.8f}\n"
        f"surface_threshold_time_h={surface_hit_h:.8f}\n"
        f"final_surface_moisture={final_surface:.8f}\n"
        f"tail_fraction_after_surface={late_fraction:.8f}\n"
        f"n_intervals=800\n"
        f"nfev={solution.nfev}\n"
    )
    (RESULTS / "q3_summary.txt").write_text(summary, encoding="utf-8")
    (RESULTS / "q3_summary.tex").write_text(
        "\\newcommand{\\QThreeCompletion}{" + f"{completion_s/3600.0:.2f}" + "}\n"
        "\\newcommand{\\QThreeSurfaceHit}{" + f"{surface_hit_h:.2f}" + "}\n"
        "\\newcommand{\\QThreeEndSurface}{" + f"{final_surface:.4f}" + "}\n"
        "\\newcommand{\\QThreeTailFraction}{" + f"{late_fraction*100.0:.1f}" + "}\n",
        encoding="utf-8",
    )
    print(summary, end="")


if __name__ == "__main__":
    main()
