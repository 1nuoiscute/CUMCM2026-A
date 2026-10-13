"""A题问题四第一版：水分状态驱动的收缩边界模型。

计算链：
1. 以附件2的实测半径作为移动边界，求得与半径数据同时刻的平均含水率；
2. 由 R/R0 与 Cbar/C0 识别低参数收缩响应曲线；
3. 将该响应曲线闭合回温度--水分方程，使半径由当前含水状态决定；
4. 输出全域达标时间、result4.xlsx、论文表格和必要图件。
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import csv
import json
import shutil
from project_paths import DATA_ROOT, OUTPUT_ROOT

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.optimize import curve_fit
from scipy.sparse import lil_matrix


ROOT = Path(__file__).resolve().parent
CONTEST_ROOT = ROOT.parent
INPUT_DIR = DATA_ROOT
ATTACHMENT1 = INPUT_DIR / "附件1.xlsx"
ATTACHMENT2 = INPUT_DIR / "附件2.xlsx"
TEMPLATE = INPUT_DIR / "附件3" / "result4.xlsx"
OUTPUT = OUTPUT_ROOT / "q4" / "result4.xlsx"
RESULTS = OUTPUT_ROOT / "q4"
FIGURES = RESULTS
PUBLISHED_FIGURES = {
    "raw_q4_radius_reconstruction",
    "process_q4_shrinkage_response",
    "result_q4_moisture_history",
}

R0 = 0.020
H = 25.0
HM = 8.0e-7
T0_K = 301.15
C0 = 2.55
C_TARGET = 0.15
PLATEAU_START_S = 4.0 * 3600.0
PLATEAU_T_C = 50.0
PLATEAU_C = 0.0500


class FullProcessAmbient:
    """附件1的0--4 h动态边界，之后采用稳定阶段均值。"""

    def __init__(self, path: Path):
        ws = load_workbook(path, data_only=True, read_only=True).active
        rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
        values = np.asarray(rows, dtype=float)
        self.time_s = values[:, 0]
        self.temperature_c = values[:, 1]
        self.moisture = values[:, 2]

    def __call__(self, time_s: float) -> tuple[float, float]:
        if time_s <= PLATEAU_START_S:
            return (
                float(np.interp(time_s, self.time_s, self.temperature_c)),
                float(np.interp(time_s, self.time_s, self.moisture)),
            )
        return PLATEAU_T_C, PLATEAU_C


class RadiusObservation:
    """附件2半径数据及其保形插值。"""

    def __init__(self, path: Path):
        ws = load_workbook(path, data_only=True, read_only=True).active
        rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
        values = np.asarray(rows, dtype=float)
        self.time_s = values[:, 0]
        self.radius_m = values[:, 1] / 100.0
        self.interpolator = PchipInterpolator(self.time_s, self.radius_m, extrapolate=False)

    def __call__(self, time_s: float | np.ndarray) -> np.ndarray:
        t = np.asarray(time_s, dtype=float)
        clipped = np.clip(t, self.time_s[0], self.time_s[-1])
        value = np.asarray(self.interpolator(clipped), dtype=float)
        value = np.where(t > self.time_s[-1], self.radius_m[-1], value)
        return value


def rho(C, T):
    return 760.0 + 90.0 * C


def cp(C, T):
    return 1850.0 + 2150.0 * C / (C + 1.0)


def conductivity(C, T):
    return 0.12 + 0.20 * C / (C + 1.0)


def diffusivity_q4(C, T):
    return 4.2e-4 * np.exp(-0.30 / np.maximum(C, 1.0e-8)) * np.exp(-3850.0 / T)


def diffusivity_q3(C, T):
    return 2.4e-3 * np.exp(-0.45 / np.maximum(C, 1.0e-8)) * np.exp(-3850.0 / T)


def shrinkage_ratio(x, a, b):
    """R/R0 = 1-a[1-(Cbar/C0)^b]，0<a<1, b>0。"""
    x = np.clip(np.asarray(x, dtype=float), 0.0, 1.0)
    return 1.0 - a * (1.0 - x**b)


def linear_shrinkage_ratio(x, a):
    x = np.asarray(x, dtype=float)
    return 1.0 - a * (1.0 - x)


def cubic_shrinkage_ratio(x, a, b, c):
    x = np.asarray(x, dtype=float)
    z = 1.0 - x
    return 1.0 - a * z - b * z**2 - c * z**3


def fixed_xi_grid(n_intervals: int):
    xi = np.linspace(0.0, 1.0, n_intervals + 1)
    dxi = 1.0 / n_intervals
    faces = np.empty(n_intervals + 2)
    faces[0] = 0.0
    faces[1:-1] = 0.5 * (xi[:-1] + xi[1:])
    faces[-1] = 1.0
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return xi, faces, volumes, dxi


def harmonic_mean(left, right):
    return 2.0 * left * right / np.maximum(left + right, np.finfo(float).tiny)


def jacobian_pattern(n_nodes: int):
    pattern = lil_matrix((2 * n_nodes, 2 * n_nodes), dtype=int)
    for eq_block in (0, 1):
        for i in range(n_nodes):
            row = eq_block * n_nodes + i
            for var_block in (0, 1):
                for j in range(max(0, i - 1), min(n_nodes, i + 2)):
                    pattern[row, var_block * n_nodes + j] = 1
    return pattern.tocsr()


@dataclass
class MovingResult:
    time_s: np.ndarray
    xi: np.ndarray
    temperature_k: np.ndarray
    moisture: np.ndarray
    radius_m: np.ndarray
    nfev: int

    @property
    def temperature_c(self):
        return self.temperature_k - 273.15


def solve_moving(
    ambient: FullProcessAmbient,
    radius_function,
    end_time_s: float,
    output_times_s: np.ndarray,
    *,
    n_intervals: int = 300,
    diffusivity=diffusivity_q4,
    max_step_s: float = 30.0,
    rtol: float = 1.0e-7,
    atol: float = 1.0e-9,
):
    """在材料坐标 xi=r/R(t) 上求解移动边界模型。"""
    xi, faces, volumes, dxi = fixed_xi_grid(n_intervals)
    n_nodes = xi.size
    weights = volumes / volumes.sum()

    def current_radius(time_s, moisture):
        mean_c = float(weights @ moisture)
        return float(np.asarray(radius_function(time_s, mean_c)))

    def rhs(time_s, state):
        T = state[:n_nodes]
        C = state[n_nodes:]
        R = current_radius(time_s, C)
        Tinf_c, Cinf = ambient(time_s)
        Tinf = Tinf_c + 273.15

        rhov = rho(C, T)
        cpv = cp(C, T)
        kv = conductivity(C, T)
        Dv = diffusivity(C, T)
        kf = harmonic_mean(kv[:-1], kv[1:])
        Df = harmonic_mean(Dv[:-1], Dv[1:])

        heat_coordinate = np.zeros(n_nodes + 1)
        mass_coordinate = np.zeros(n_nodes + 1)
        heat_coordinate[1:-1] = faces[1:-1] * kf * np.diff(T) / dxi
        mass_coordinate[1:-1] = faces[1:-1] * Df * np.diff(C) / dxi
        heat_coordinate[-1] = R * H * (Tinf - T[-1])
        mass_coordinate[-1] = R * HM * (Cinf - C[-1])

        heat_rate = np.diff(heat_coordinate) / (volumes * R * R * rhov * cpv)
        moisture_rate = np.diff(mass_coordinate) / (volumes * R * R)
        return np.concatenate([heat_rate, moisture_rate])

    initial = np.concatenate([np.full(n_nodes, T0_K), np.full(n_nodes, C0)])
    sol = solve_ivp(
        rhs,
        (0.0, float(end_time_s)),
        initial,
        method="BDF",
        t_eval=np.asarray(output_times_s, dtype=float),
        max_step=max_step_s,
        rtol=rtol,
        atol=atol,
        jac_sparsity=jacobian_pattern(n_nodes),
    )
    if not sol.success:
        raise RuntimeError(sol.message)
    moisture = sol.y[n_nodes:]
    radius = np.asarray([
        current_radius(t, moisture[:, j]) for j, t in enumerate(sol.t)
    ])
    return MovingResult(sol.t, xi, sol.y[:n_nodes], moisture, radius, sol.nfev)


def mean_field(result: MovingResult, field: np.ndarray):
    _, _, volumes, _ = fixed_xi_grid(result.xi.size - 1)
    return (volumes / volumes.sum()) @ field


def field_at_time(time_s, field, target_s):
    i1 = int(np.searchsorted(time_s, target_s, side="left"))
    if i1 == 0:
        return field[:, 0].copy()
    if i1 >= time_s.size:
        return field[:, -1].copy()
    if time_s[i1] == target_s:
        return field[:, i1].copy()
    i0 = i1 - 1
    w = (target_s - time_s[i0]) / (time_s[i1] - time_s[i0])
    return (1.0 - w) * field[:, i0] + w * field[:, i1]


def scalar_at_time(time_s, values, target_s):
    return float(np.interp(target_s, time_s, values))


def completion_time(result: MovingResult):
    maximum = result.moisture.max(axis=0)
    hits = np.flatnonzero(maximum <= C_TARGET)
    if hits.size == 0:
        raise RuntimeError("计算时间内尚未全域达标")
    i1 = int(hits[0])
    i0 = i1 - 1
    w = (maximum[i0] - C_TARGET) / (maximum[i0] - maximum[i1])
    return float(result.time_s[i0] + w * (result.time_s[i1] - result.time_s[i0])), i1


def sample_physical_profile(result: MovingResult, target_s: float, radii_m):
    C = field_at_time(result.time_s, result.moisture, target_s)
    R = scalar_at_time(result.time_s, result.radius_m, target_s)
    radii = np.asarray(radii_m, dtype=float)
    values = np.full(radii.shape, np.nan)
    valid = radii <= R + 1.0e-12
    values[valid] = np.interp(radii[valid] / R, result.xi, C)
    return values, R, C[-1]


def identify_shrinkage_law(ambient, radius_obs: RadiusObservation):
    """由附件2约束的移动边界解识别 R/R0--Cbar/C0 收缩响应。"""
    calibration_times = radius_obs.time_s
    observed_model = solve_moving(
        ambient,
        lambda t, cbar: radius_obs(t),
        radius_obs.time_s[-1],
        calibration_times,
        n_intervals=240,
        max_step_s=60.0,
        rtol=3.0e-7,
        atol=3.0e-9,
    )
    cbar = mean_field(observed_model, observed_model.moisture)
    x = cbar / C0
    y = radius_obs.radius_m / R0
    train = calibration_times <= 48.0 * 3600.0
    params, _ = curve_fit(
        shrinkage_ratio,
        x[train],
        y[train],
        p0=(0.42, 1.0),
        bounds=([0.05, 0.05], [0.80, 8.0]),
        maxfev=20000,
    )
    predicted = shrinkage_ratio(x, *params)
    train_rmse = float(np.sqrt(np.mean((predicted[train] - y[train]) ** 2)))
    test_rmse = float(np.sqrt(np.mean((predicted[~train] - y[~train]) ** 2)))
    return params, observed_model, x, y, predicted, train_rmse, test_rmse


def compare_shrinkage_candidates(x, y, train_mask):
    """按统一样本和厘米口径比较收缩响应候选函数。"""
    candidates = [
        ("线性", r"$y=1-a(1-x)$", linear_shrinkage_ratio,
         (0.4,), ([0.0], [1.0])),
        ("幂函数", r"$y=1-a(1-x^b)$", shrinkage_ratio,
         (0.42, 1.0), ([0.05, 0.05], [0.80, 8.0])),
        ("三次多项式", r"三次多项式（约束 $y(1)=1$）", cubic_shrinkage_ratio,
         (0.4, 0.0, 0.0), ([-5.0, -5.0, -5.0], [5.0, 5.0, 5.0])),
    ]
    rows = []
    radius_scale_cm = R0 * 100.0
    for name, latex_name, function, initial, bounds in candidates:
        params, _ = curve_fit(
            function,
            x[train_mask],
            y[train_mask],
            p0=initial,
            bounds=bounds,
            maxfev=20000,
        )
        predicted = function(x, *params)
        train_rmse_cm = float(
            np.sqrt(np.mean((predicted[train_mask] - y[train_mask]) ** 2))
            * radius_scale_cm
        )
        test_rmse_cm = float(
            np.sqrt(np.mean((predicted[~train_mask] - y[~train_mask]) ** 2))
            * radius_scale_cm
        )
        rows.append((name, latex_name, len(params), train_rmse_cm, test_rmse_cm))
    return rows


def write_shrinkage_candidate_results(rows):
    with (RESULTS / "q4_shrinkage_candidates.csv").open(
        "w", newline="", encoding="utf-8-sig"
    ) as file:
        writer = csv.writer(file)
        writer.writerow(["候选模型", "参数数", "前48h_RMSE_cm", "48-72h_RMSE_cm"])
        for name, _latex_name, n_params, train_rmse, test_rmse in rows:
            writer.writerow([name, n_params, f"{train_rmse:.8f}", f"{test_rmse:.8f}"])

    latex_rows = []
    for index, (_name, latex_name, n_params, train_rmse, test_rmse) in enumerate(rows):
        row_ending = r"\\" if index < len(rows) - 1 else ""
        latex_rows.append(
            f"{latex_name} & {n_params} & {train_rmse:.4f} & {test_rmse:.4f}{row_ending}"
        )
    (RESULTS / "q4_shrinkage_candidate_rows.tex").write_text(
        "\n".join(latex_rows) + "\n", encoding="utf-8"
    )


def setup_plot():
    available = {f.name for f in mpl.font_manager.fontManager.ttflist}
    font = "Microsoft YaHei" if "Microsoft YaHei" in available else "SimHei"
    mpl.rcParams.update({
        "font.sans-serif": [font],
        "axes.unicode_minus": False,
        "figure.dpi": 130,
        "savefig.dpi": 320,
        "axes.spines.top": False,
        "axes.spines.right": False,
        "font.size": 9,
    })


def save_figure(fig, name):
    if name in PUBLISHED_FIGURES:
        fig.savefig(FIGURES / f"{name}.png", dpi=320, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_figures(result, radius_obs, law_data, completion_s, fixed_q4_time_h, q3_time_h):
    setup_plot()
    params, _obs_model, x, y, yfit, train_rmse, test_rmse = law_data
    time_h = result.time_s / 3600.0
    cbar = mean_field(result, result.moisture)

    # 1. 附件半径与闭环模型预测
    predicted_at_obs = np.interp(radius_obs.time_s, result.time_s, result.radius_m)
    fig, ax = plt.subplots(figsize=(6.5, 3.7), constrained_layout=True)
    ax.scatter(radius_obs.time_s / 3600.0, radius_obs.radius_m * 100.0,
               s=13, color="#264653", label="附件2实测半径", zorder=3)
    ax.plot(radius_obs.time_s / 3600.0, predicted_at_obs * 100.0,
            color="#E76F51", lw=1.6, label="收缩响应模型")
    ax.axvline(48.0, color="#777777", lw=1.0, ls="--", label="辨识/检验分界")
    ax.set(xlabel="时间 / h", ylabel="药材半径 / cm", xlim=(0, 72))
    ax.legend(frameon=False)
    save_figure(fig, "raw_q4_radius_reconstruction")

    # 2. 收缩响应曲线
    order = np.argsort(x)
    xx = np.linspace(max(0.0, x.min()), 1.0, 300)
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.25), constrained_layout=True)
    axes[0].scatter(x, y, s=12, color="#2A9D8F", alpha=0.75, label="附件2映射点")
    axes[0].plot(xx, shrinkage_ratio(xx, *params), color="#D1495B", lw=1.6,
                 label="两参数响应曲线")
    axes[0].set(xlabel=r"平均含水率比 $\bar C/C_0$", ylabel=r"半径比 $R/R_0$")
    axes[0].legend(frameon=False, fontsize=8)
    response = params[0] * params[1] * np.maximum(xx, 1e-6) ** (params[1] - 1.0)
    axes[1].plot(xx, response, color="#6A994E", lw=1.6)
    axes[1].set(xlabel=r"平均含水率比 $\bar C/C_0$",
                ylabel=r"结构响应 $d(R/R_0)/d(\bar C/C_0)$")
    save_figure(fig, "process_q4_shrinkage_response")

    # 3. 水分历史
    fig, ax = plt.subplots(figsize=(6.6, 3.8), constrained_layout=True)
    for xi_value, color, label in [
        (0.0, "#264653", "中心"),
        (0.5, "#2A9D8F", r"材料位置 $\xi=0.5$"),
        (0.75, "#E9C46A", r"材料位置 $\xi=0.75$"),
        (1.0, "#E76F51", "表面"),
    ]:
        values = np.vstack([
            np.interp(xi_value, result.xi, result.moisture[:, j])
            for j in range(result.time_s.size)
        ]).ravel()
        ax.plot(time_h, values, color=color, lw=1.35, label=label)
    ax.plot(time_h, cbar, color="#555555", lw=1.2, ls="--", label="截面平均")
    ax.axhline(C_TARGET, color="#B22222", lw=1.0, ls=":")
    ax.axvline(completion_s / 3600.0, color="#B22222", lw=1.0, ls="--")
    ax.set(xlabel="时间 / h", ylabel="干基含水率 / (kg/kg)",
           xlim=(0, completion_s / 3600.0))
    ax.legend(frameon=False, ncol=3, fontsize=8)
    save_figure(fig, "result_q4_moisture_history")

    # 4. 物理坐标下的剖面
    profile_hours = [6, 12, 24, 36, 48, completion_s / 3600.0]
    colors = ["#440154", "#3B528B", "#21918C", "#5EC962", "#FDE725", "#D1495B"]
    fig, ax = plt.subplots(figsize=(6.4, 3.75), constrained_layout=True)
    for hour, color in zip(profile_hours, colors):
        C = field_at_time(result.time_s, result.moisture, hour * 3600.0)
        R = scalar_at_time(result.time_s, result.radius_m, hour * 3600.0)
        label = f"{hour:.2f} h" if hour == profile_hours[-1] else f"{hour:g} h"
        ax.plot(result.xi * R * 100.0, C, color=color, lw=1.5, label=label)
    ax.axhline(C_TARGET, color="#B22222", lw=1.0, ls=":")
    ax.set(xlabel="到药材中心的实际距离 / cm", ylabel="干基含水率 / (kg/kg)")
    ax.legend(frameon=False, ncol=3, fontsize=8)
    save_figure(fig, "result_q4_moving_profiles")

    # 5. 烘干时间机制比较
    fig, ax = plt.subplots(figsize=(5.6, 3.5), constrained_layout=True)
    names = ["问题三\n固定半径+附录3", "仅换附录4\n固定半径", "问题四\n附录4+收缩"]
    values = [q3_time_h, fixed_q4_time_h, completion_s / 3600.0]
    bars = ax.bar(names, values, color=["#8DA0CB", "#66C2A5", "#FC8D62"], width=0.62)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.8, f"{value:.2f} h",
                ha="center", va="bottom", fontsize=8)
    ax.set(ylabel="全域达标时间 / h", ylim=(0, max(values) * 1.16))
    save_figure(fig, "result_q4_mechanism_comparison")


def save_workbook(result: MovingResult, completion_index: int):
    physical_cm = np.arange(0.0, 2.0 + 0.1, 0.1)
    shutil.copy2(TEMPLATE, OUTPUT)
    wb = load_workbook(OUTPUT)
    ws = wb.active
    ws.delete_rows(1, ws.max_row)
    ws.append(["时间\\到药材中心的距离", *[round(float(x), 1) for x in physical_cm], "药材表面"])
    for j in range(1, completion_index + 1):
        R = result.radius_m[j]
        C = result.moisture[:, j]
        row = [int(round(result.time_s[j]))]
        for radius_cm in physical_cm:
            radius_m = radius_cm / 100.0
            if radius_m <= R + 1.0e-12:
                row.append(round(float(np.interp(radius_m / R, result.xi, C)), 4))
            else:
                row.append(None)
        row.append(round(float(C[-1]), 4))
        ws.append(row)
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 24
    for column in ws.iter_cols(min_col=2, max_col=ws.max_column):
        ws.column_dimensions[column[0].column_letter].width = 10
    wb.save(OUTPUT)


def make_table(result: MovingResult, completion_s: float):
    hours = np.arange(6.0, np.floor(completion_s / 21600.0) * 6.0 + 0.1, 6.0)
    times = np.append(hours, completion_s / 3600.0)
    lines = []
    csv_rows = []
    for index, hour in enumerate(times):
        values, R, Cs = sample_physical_profile(
            result, hour * 3600.0, np.array([0.0, 0.005, 0.010])
        )
        name = f"{int(hour)}" if index < len(times) - 1 else f"烘干结束（{hour:.2f} h）"
        fields = [name, f"{R*100:.4f}", *[f"{v:.4f}" for v in values], f"{Cs:.4f}"]
        lines.append(" & ".join(fields) + r" \\")
        csv_rows.append([hour, R * 100.0, *values, Cs])
    (RESULTS / "q4_table_rows.tex").write_text(
        "\n".join(lines) + "\n\\bottomrule\n", encoding="utf-8"
    )
    np.savetxt(
        RESULTS / "q4_table.csv",
        np.asarray(csv_rows),
        delimiter=",",
        header="time_h,radius_cm,C_0cm,C_0p5cm,C_1cm,C_surface",
        comments="",
        fmt="%.8f",
    )


def q3_completion_from_existing():
    path = OUTPUT_ROOT / "q3" / "q3_summary.txt"
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.startswith("completion_time_h="):
                return float(line.split("=", 1)[1])
    return 57.51225261


def main():
    RESULTS.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    ambient = FullProcessAmbient(ATTACHMENT1)
    radius_obs = RadiusObservation(ATTACHMENT2)

    params, observed_model, x, y, yfit, train_rmse, test_rmse = identify_shrinkage_law(
        ambient, radius_obs
    )
    a, b = params
    train_mask = radius_obs.time_s <= 48.0 * 3600.0
    candidate_rows = compare_shrinkage_candidates(x, y, train_mask)
    write_shrinkage_candidate_results(candidate_rows)

    # 闭环模型：当前平均含水率决定当前半径。
    response_radius = lambda t, cbar: R0 * shrinkage_ratio(cbar / C0, a, b)
    output_times = np.arange(0.0, 96.0 * 3600.0 + 60.0, 60.0)
    result = solve_moving(
        ambient,
        response_radius,
        output_times[-1],
        output_times,
        n_intervals=300,
        max_step_s=30.0,
        rtol=1.0e-7,
        atol=1.0e-9,
    )
    completion_s, completion_index = completion_time(result)

    # 固定半径+附录4用于拆分“物性改变”和“收缩”两种作用。
    comparison_times = np.arange(0.0, 240.0 * 3600.0 + 300.0, 300.0)
    fixed_q4 = solve_moving(
        ambient,
        lambda t, cbar: R0,
        comparison_times[-1],
        comparison_times,
        n_intervals=180,
        max_step_s=60.0,
        rtol=3.0e-7,
        atol=3.0e-9,
    )
    fixed_q4_s, _ = completion_time(fixed_q4)
    fixed_q4_h = fixed_q4_s / 3600.0
    q3_h = q3_completion_from_existing()

    save_workbook(result, completion_index)
    make_table(result, completion_s)
    make_figures(
        result,
        radius_obs,
        (params, observed_model, x, y, yfit, train_rmse, test_rmse),
        completion_s,
        fixed_q4_h,
        q3_h,
    )

    predicted_radius_obs = np.interp(radius_obs.time_s, result.time_s, result.radius_m)
    radius_rmse_cm = float(np.sqrt(np.mean((predicted_radius_obs-radius_obs.radius_m)**2))*100)
    test_mask = radius_obs.time_s > 48.0 * 3600.0
    radius_test_rmse_cm = float(
        np.sqrt(np.mean((predicted_radius_obs[test_mask]-radius_obs.radius_m[test_mask])**2))*100
    )
    final_profile = field_at_time(result.time_s, result.moisture, completion_s)
    final_radius = scalar_at_time(result.time_s, result.radius_m, completion_s)
    shrink_advance_h = fixed_q4_h - completion_s / 3600.0
    property_delay_h = fixed_q4_h - q3_h

    summary = (
        f"completion_time_h={completion_s/3600.0:.8f}\n"
        f"first_output_completion_h={result.time_s[completion_index]/3600.0:.8f}\n"
        f"radius_at_completion_cm={final_radius*100.0:.8f}\n"
        f"surface_moisture_at_completion={final_profile[-1]:.8f}\n"
        f"shrinkage_a={a:.8f}\n"
        f"shrinkage_b={b:.8f}\n"
        f"calibration_curve_train_rmse={train_rmse:.8e}\n"
        f"calibration_curve_test_rmse={test_rmse:.8e}\n"
        f"coupled_radius_rmse_cm={radius_rmse_cm:.8f}\n"
        f"coupled_radius_test_rmse_cm={radius_test_rmse_cm:.8f}\n"
        f"candidate_total_points={radius_obs.time_s.size}\n"
        f"candidate_train_points={int(train_mask.sum())}\n"
        f"candidate_test_points={int((~train_mask).sum())}\n"
        f"fixed_radius_appendix4_completion_h={fixed_q4_h:.8f}\n"
        f"q3_completion_h={q3_h:.8f}\n"
        f"property_change_delay_h={property_delay_h:.8f}\n"
        f"shrinkage_time_reduction_h={shrink_advance_h:.8f}\n"
        f"n_intervals=300\n"
        f"nfev={result.nfev}\n"
    )
    (RESULTS / "q4_summary.txt").write_text(summary, encoding="utf-8")
    (RESULTS / "q4_summary.tex").write_text(
        "\\newcommand{\\QFourCompletion}{" + f"{completion_s/3600.0:.2f}" + "}\n"
        "\\newcommand{\\QFourRadiusEnd}{" + f"{final_radius*100.0:.3f}" + "}\n"
        "\\newcommand{\\QFourSurfaceEnd}{" + f"{final_profile[-1]:.4f}" + "}\n"
        "\\newcommand{\\QFourShrinkA}{" + f"{a:.4f}" + "}\n"
        "\\newcommand{\\QFourShrinkB}{" + f"{b:.4f}" + "}\n"
        "\\newcommand{\\QFourRadiusRMSE}{" + f"{radius_rmse_cm:.4f}" + "}\n"
        "\\newcommand{\\QFourRadiusTestRMSE}{" + f"{radius_test_rmse_cm:.4f}" + "}\n"
        "\\newcommand{\\QFourCandidateTotal}{" + f"{radius_obs.time_s.size}" + "}\n"
        "\\newcommand{\\QFourCandidateTrain}{" + f"{int(train_mask.sum())}" + "}\n"
        "\\newcommand{\\QFourCandidateTest}{" + f"{int((~train_mask).sum())}" + "}\n"
        "\\newcommand{\\QFourFixedTime}{" + f"{fixed_q4_h:.2f}" + "}\n"
        "\\newcommand{\\QFourPropertyDelay}{" + f"{property_delay_h:.2f}" + "}\n"
        "\\newcommand{\\QFourShrinkGain}{" + f"{shrink_advance_h:.2f}" + "}\n",
        encoding="utf-8",
    )
    manifest = {
        "command": "python 问题4.py",
        "model": "含水状态驱动的收缩边界模型",
        "inputs": [str(ATTACHMENT1), str(ATTACHMENT2), str(TEMPLATE)],
        "outputs": [str(OUTPUT), str(RESULTS), str(FIGURES)],
        "n_intervals": 300,
        "max_step_s": 30.0,
        "random_seed": None,
    }
    (RESULTS / "复现清单.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(summary, end="")


if __name__ == "__main__":
    main()
